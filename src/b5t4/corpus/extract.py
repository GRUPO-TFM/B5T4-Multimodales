"""
Extracción de contenido de los ficheros del corpus, por tipo.

Devuelve «páginas» homogéneas (texto + de dónde sale) para que los agentes no tengan que saber si algo
venía de un PDF, de un HTML o de un correo. Para los PDF puede además rasterizar páginas a PNG: es lo que
necesita el agente 2 para que un modelo de visión lea gráficos y tablas que no tienen texto extraíble.

Tipos cubiertos: .md .txt .csv .json · .html · .eml · .pdf · .pptx.
Los .xlsx (maestros con fórmulas Bloomberg) quedan fuera a propósito: sin Terminal sus valores no son fiables.
"""
from __future__ import annotations

import email
import email.policy
import html as html_lib
import re
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

TEXT_EXTS = {".md", ".txt", ".csv", ".json"}
SUPPORTED = TEXT_EXTS | {".html", ".htm", ".eml", ".pdf", ".pptx"}


@dataclass
class Page:
    source: str      # ruta relativa del fichero
    page: int        # 1-based; 1 para documentos sin páginas
    text: str


class _HTMLText(HTMLParser):
    """HTML → texto plano conservando saltos de bloque. Sin dependencias externas."""
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "table", "td", "th"}

    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self._skip = max(0, self._skip - 1)
        elif tag in ("td", "th"):
            self.parts.append(" | ")

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def html_to_text(raw: str) -> str:
    p = _HTMLText()
    p.feed(raw)
    txt = html_lib.unescape("".join(p.parts))
    txt = re.sub(r"[ \t\xa0]+", " ", txt)
    return re.sub(r"\n\s*\n+", "\n\n", txt).strip()


def _eml_text(path: Path) -> str:
    msg = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
    head = f"Asunto: {msg.get('Subject', '')}\nPara: {msg.get('To', '')}\nCc: {msg.get('Cc', '')}\n\n"
    body = msg.get_body(preferencelist=("plain", "html"))
    if body is None:
        return head
    content = body.get_content()
    if body.get_content_type() == "text/html" or content.lstrip().startswith("<"):
        content = html_to_text(content)
    if content.strip().lower().startswith("ver versión html"):   # nuestros .eml llevan el HTML como alternativa
        html_part = msg.get_body(preferencelist=("html",))
        content = html_to_text(html_part.get_content()) if html_part else content
    return head + content


def pdf_pages(path: Path, source: str = "") -> list[Page]:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        return [Page(source or path.name, i + 1, pdf[i].get_textpage().get_text_range()) for i in range(len(pdf))]
    finally:
        pdf.close()


def pdf_page_images(path: Path, out_dir: Path, pages: list[int] | None = None, scale: float = 2.0) -> list[Path]:
    """Rasteriza páginas (1-based) a PNG para un modelo de visión. scale=2 ≈ 144 ppp: suficiente para leer tablas."""
    import pypdfium2 as pdfium

    out_dir.mkdir(parents=True, exist_ok=True)
    pdf = pdfium.PdfDocument(str(path))
    try:
        idx = pages or list(range(1, len(pdf) + 1))
        outs = []
        for n in idx:
            img = pdf[n - 1].render(scale=scale).to_pil()
            p = out_dir / f"{path.stem}_p{n:02d}.png"
            img.save(p)
            outs.append(p)
        return outs
    finally:
        pdf.close()


def _pptx_pages(path: Path, source: str) -> list[Page]:
    from pptx import Presentation

    pages = []
    for i, slide in enumerate(Presentation(str(path)).slides, start=1):
        txt = [sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
        pages.append(Page(source, i, "\n".join(txt)))
    return pages


def extract(path: Path, source: str = "") -> list[Page]:
    """Punto de entrada único: fichero → lista de páginas de texto. Lista vacía si el tipo no se soporta."""
    source = source or path.name
    ext = path.suffix.lower()
    if ext in TEXT_EXTS:
        return [Page(source, 1, path.read_text(encoding="utf-8", errors="replace"))]
    if ext in (".html", ".htm"):
        return [Page(source, 1, html_to_text(path.read_text(encoding="utf-8", errors="replace")))]
    if ext == ".eml":
        return [Page(source, 1, _eml_text(path))]
    if ext == ".pdf":
        return pdf_pages(path, source)
    if ext == ".pptx":
        return _pptx_pages(path, source)
    return []
